"""Build the document fixtures used by the extraction tests.

Generated rather than committed as binaries: a .docx or .xlsx in Git is an
opaque blob that nobody can review, and regenerating them keeps the fixtures
honest about which library version produced them.
"""

from __future__ import annotations

from pathlib import Path

RESEARCH_TEXT = (
    "Tactile feedback improves grasp success on deformable objects. "
    "Across 240 trials the tactile-enabled policy reached 94 percent success, "
    "against 83 percent for the vision-only baseline. "
    "The remaining failures were dominated by slip during transport rather than "
    "by initial grasp selection. "
    "Future work should extend the evaluation to objects outside the training "
    "distribution, particularly thin and highly compliant items."
)


def write_text_fixtures(directory: Path) -> None:
    (directory / "research_notes.txt").write_text(RESEARCH_TEXT, encoding="utf-8")
    (directory / "meeting.md").write_text(
        "# Lab meeting\n\n"
        "## Decisions\n\n"
        "- Rerun the ablation with tactile disabled.\n"
        "- Sam to draft the related work section by Friday.\n\n"
        "The grasp benchmark is now the primary evaluation.\n",
        encoding="utf-8",
    )
    (directory / "config.yaml").write_text(
        "training:\n  epochs: 40\n  batch_size: 32\n  seed: 7\n", encoding="utf-8"
    )
    (directory / "train.py").write_text(
        '"""Training entry point."""\n\n\ndef main() -> None:\n    print("training")\n',
        encoding="utf-8",
    )
    (directory / "results.csv").write_text(
        "trial,object,success,slip\n1,sponge,1,0\n2,bottle,1,0\n3,cloth,0,1\n",
        encoding="utf-8",
    )
    (directory / "run.json").write_text(
        '{"run_id": "abc123", "epochs": 40, "final_loss": 0.031, '
        '"notes": "tactile ablation complete"}',
        encoding="utf-8",
    )
    (directory / "ignored.bin").write_bytes(b"\x00\x01\x02binary payload")


def write_pdf(path: Path) -> None:
    from pypdf import PdfWriter

    # pypdf cannot author content streams, so the PDF is assembled by hand. It is
    # the minimum structure a reader needs: one page, one font, one text object.
    text = "Tactile feedback improves grasp success on deformable objects."
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")

    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    ).encode()

    path.write_bytes(bytes(out))
    PdfWriter(clone_from=str(path))  # fails loudly if the structure is wrong


def write_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Quarterly research summary", level=1)
    document.add_paragraph(RESEARCH_TEXT)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Policy"
    table.cell(0, 1).text = "Success"
    table.cell(1, 0).text = "Tactile"
    table.cell(1, 1).text = "94%"
    document.save(str(path))


def write_pptx(path: Path) -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Grasp benchmark results"
    slide.placeholders[1].text = "Tactile policy reaches 94 percent success"
    presentation.save(str(path))


def write_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Trials"
    sheet.append(["trial", "object", "success"])
    sheet.append([1, "sponge", 1])
    sheet.append([2, "bottle", 1])
    workbook.save(str(path))


def build_all(directory: Path) -> Path:
    """Create one file of every supported format and return the directory."""
    directory.mkdir(parents=True, exist_ok=True)
    write_text_fixtures(directory)
    write_pdf(directory / "paper.pdf")
    write_docx(directory / "summary.docx")
    write_pptx(directory / "slides.pptx")
    write_xlsx(directory / "trials.xlsx")
    return directory


if __name__ == "__main__":
    import sys

    target = Path(sys.argv[1] if len(sys.argv) > 1 else "./document_fixtures")
    print(f"Wrote fixtures to {build_all(target)}")
